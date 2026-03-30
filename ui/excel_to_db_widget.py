"""
Excel数据导入数据库功能界面
"""
import os
import pandas as pd
from datetime import datetime
from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QFormLayout,
                             QPushButton, QLabel, QFileDialog, QMessageBox,
                             QTableWidget, QTableWidgetItem, QComboBox,
                             QGroupBox, QProgressBar, QLineEdit, QTextEdit,
                             QCheckBox, QSpinBox, QHeaderView)
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtGui import QFont
from database import DatabaseManager
import pymysql


class ExcelToDbWorker(QThread):
    """Excel导入数据库工作线程"""
    progress = pyqtSignal(int)
    log_message = pyqtSignal(str)
    finished = pyqtSignal(dict)
    error = pyqtSignal(str)
    
    # import_mode: 'non_corporate' = 非企业法人(is_customer=1), 'corporate' = 企业法人(is_customer=2)
    def __init__(self, input_file, datasource, db_manager, column_mapping, check_field, import_mode):
        super().__init__()
        self.input_file = input_file
        self.datasource = datasource
        self.db_manager = db_manager
        self.column_mapping = column_mapping
        self.check_field = check_field
        self.import_mode = import_mode  # 'non_corporate' or 'corporate'
        self.results = {
            'total': 0,
            'success': 0,
            'skipped_empty': 0,
            'skipped_exists': 0,
            'skipped_wrong_type': 0,
            'failed': 0,
            'logs': []
        }
    
    def run(self):
        try:
            mode_name = "非企业法人" if self.import_mode == 'non_corporate' else "企业法人"
            is_customer_value = 1 if self.import_mode == 'non_corporate' else 2
            
            self.log_message.emit(f"开始读取Excel文件... (模式: {mode_name}, is_customer={is_customer_value})")
            self.progress.emit(5)
            
            if self.input_file.endswith('.xlsx'):
                try:
                    df = pd.read_excel(self.input_file, engine='openpyxl', header=0)
                except:
                    from openpyxl import load_workbook
                    wb = load_workbook(self.input_file, data_only=True)
                    ws = wb.active
                    
                    data = []
                    for row in ws.iter_rows(values_only=True):
                        data.append(row)
                    
                    if data:
                        columns = [str(col) if col is not None else f'列{i+1}' for i, col in enumerate(data[0])]
                        df = pd.DataFrame(data[1:], columns=columns)
            else:
                df = pd.read_excel(self.input_file, engine='xlrd', header=0)
            
            self.results['total'] = len(df)
            self.log_message.emit(f"读取到 {len(df)} 行数据")
            self.progress.emit(10)
            
            self.log_message.emit("连接数据库...")
            connection = pymysql.connect(
                host=self.datasource.host,
                port=self.datasource.port,
                user=self.datasource.username,
                password=self.datasource.password,
                database=self.datasource.database,
                charset=self.datasource.charset
            )
            
            cursor = connection.cursor()
            self.progress.emit(15)
            
            for index, row in df.iterrows():
                try:
                    progress = 15 + int((index / len(df)) * 80)
                    self.progress.emit(progress)
                    
                    # 获取客户类型，根据模式筛选
                    customer_type_col = self.column_mapping.get('customer_type')
                    customer_type_value = None
                    if customer_type_col is not None and customer_type_col >= 0:
                        customer_type_value = self.get_cell_value(row, customer_type_col)
                    
                    # 根据导入模式筛选
                    if self.import_mode == 'non_corporate':
                        # 非企业法人模式：跳过企业法人
                        if customer_type_value == "企业法人":
                            self.results['skipped_wrong_type'] += 1
                            continue
                    else:
                        # 企业法人模式：只处理企业法人
                        if customer_type_value != "企业法人":
                            self.results['skipped_wrong_type'] += 1
                            continue
                    
                    # 获取法人姓名
                    legal_name_col = self.column_mapping.get('legal_name')
                    if legal_name_col is None or legal_name_col < 0:
                        self.log_message.emit(f"第{index+2}行: 未配置法人姓名列，跳过")
                        self.results['skipped_empty'] += 1
                        continue
                    
                    legal_name = self.get_cell_value(row, legal_name_col)
                    if not legal_name:
                        self.log_message.emit(f"第{index+2}行: 法人姓名为空，跳过")
                        self.results['skipped_empty'] += 1
                        continue
                    
                    # 检查是否已存在
                    if self.check_field:
                        cursor.execute(f"SELECT id FROM ba_rlb_customer WHERE {self.check_field} = %s", (legal_name,))
                        existing = cursor.fetchone()
                        
                        if existing:
                            self.log_message.emit(f"第{index+2}行: '{legal_name}' 已存在（ID={existing[0]}），跳过")
                            self.results['skipped_exists'] += 1
                            continue
                    
                    # 获取担当人员
                    rlb_staff_id = None
                    admin_id = 1
                    admin_dept_id = 1
                    
                    staff_col = self.column_mapping.get('staff')
                    if staff_col is not None and staff_col >= 0:
                        staff_value = self.get_cell_value(row, staff_col)
                        if staff_value:
                            cursor.execute("""
                                SELECT id, dept_id FROM ba_admin 
                                WHERE username = %s OR nickname = %s
                            """, (staff_value, staff_value))
                            admin_result = cursor.fetchone()
                            
                            if admin_result:
                                admin_id = admin_result[0]
                                rlb_staff_id = admin_result[0]
                                admin_dept_id = admin_result[1] or 1
                                self.log_message.emit(f"第{index+2}行: 找到担当人员 '{staff_value}' (ID: {admin_id})")
                            else:
                                self.log_message.emit(f"第{index+2}行: 未找到担当人员 '{staff_value}'，使用默认值")
                    
                    # 插入数据，使用根据模式确定的 is_customer 值
                    current_time = int(datetime.now().timestamp())
                    
                    insert_sql = """
                        INSERT INTO ba_rlb_customer 
                        (legal_name, is_customer, rlb_staff_id, admin_id, admin_dept_id, remark, create_time, update_time)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                    """
                    
                    cursor.execute(insert_sql, (
                        legal_name,
                        is_customer_value,
                        rlb_staff_id,
                        admin_id,
                        admin_dept_id,
                        "系统导入",
                        current_time,
                        current_time
                    ))
                    
                    self.results['success'] += 1
                    self.log_message.emit(f"第{index+2}行: 成功插入 '{legal_name}' (is_customer={is_customer_value})")
                    
                except Exception as e:
                    self.results['failed'] += 1
                    error_msg = f"第{index+2}行: 插入失败 - {str(e)}"
                    self.log_message.emit(error_msg)
                    self.results['logs'].append(error_msg)
            
            connection.commit()
            cursor.close()
            connection.close()
            
            self.progress.emit(100)
            self.finished.emit(self.results)
            
        except Exception as e:
            self.error.emit(str(e))
    
    def get_cell_value(self, row, col_index):
        """获取单元格值"""
        if col_index is None or col_index < 0 or col_index >= len(row):
            return None
        value = row.iloc[col_index]
        if pd.isna(value) or str(value).strip() == '' or str(value).strip().lower() == 'nan':
            return None
        return str(value).strip()


class ExcelToDbWidget(QWidget):
    """Excel导入数据库功能组件"""
    
    def __init__(self, db_manager: DatabaseManager):
        super().__init__()
        self.db_manager = db_manager
        self.input_file = ""
        self.df = None
        self.init_ui()
        self.load_datasources()
    
    def init_ui(self):
        """初始化UI"""
        layout = QVBoxLayout()
        
        # 数据源选择区域
        datasource_group = QGroupBox("选择目标数据库")
        datasource_layout = QFormLayout()
        
        self.datasource_combo = QComboBox()
        self.datasource_combo.currentTextChanged.connect(self.on_datasource_changed)
        datasource_layout.addRow("数据源:", self.datasource_combo)
        
        self.test_connection_btn = QPushButton("测试连接")
        self.test_connection_btn.clicked.connect(self.test_connection)
        self.test_connection_btn.setEnabled(False)
        datasource_layout.addRow("", self.test_connection_btn)
        
        datasource_group.setLayout(datasource_layout)
        layout.addWidget(datasource_group)
        
        # 文件选择区域
        file_group = QGroupBox("选择Excel文件")
        file_layout = QFormLayout()
        
        input_layout = QHBoxLayout()
        self.input_file_label = QLabel("请选择Excel文件...")
        self.input_file_label.setStyleSheet("color: gray; font-style: italic;")
        input_layout.addWidget(self.input_file_label)
        
        self.select_file_btn = QPushButton("选择文件")
        self.select_file_btn.clicked.connect(self.select_input_file)
        input_layout.addWidget(self.select_file_btn)
        
        file_layout.addRow("Excel文件:", input_layout)
        file_group.setLayout(file_layout)
        layout.addWidget(file_group)
        
        # 列映射配置区域
        mapping_group = QGroupBox("列映射配置（选择Excel文件后可配置）")
        mapping_layout = QFormLayout()
        
        # 法人姓名列（必选）
        self.legal_name_col_combo = QComboBox()
        self.legal_name_col_combo.setEnabled(False)
        self.legal_name_col_combo.currentIndexChanged.connect(self.update_import_button_state)
        mapping_layout.addRow("法人姓名列 (必选，对应legal_name):", self.legal_name_col_combo)
        
        # 客户类型列（必选，用于区分企业法人/非企业法人）
        self.customer_type_col_combo = QComboBox()
        self.customer_type_col_combo.setEnabled(False)
        self.customer_type_col_combo.currentIndexChanged.connect(self.update_import_button_state)
        mapping_layout.addRow("客户类型列 (必选，用于区分企业法人):", self.customer_type_col_combo)
        
        # 担当人员列（可选）
        self.staff_col_combo = QComboBox()
        self.staff_col_combo.setEnabled(False)
        mapping_layout.addRow("担当人员列 (可选，匹配ba_admin表):", self.staff_col_combo)
        
        # 查重字段选择
        self.check_field_combo = QComboBox()
        self.check_field_combo.addItem("legal_name (法人姓名)", "legal_name")
        self.check_field_combo.addItem("不检查重复", "")
        mapping_layout.addRow("导入前检查重复字段:", self.check_field_combo)
        
        mapping_group.setLayout(mapping_layout)
        layout.addWidget(mapping_group)
        
        # 数据预览区域
        preview_group = QGroupBox("数据预览（前5行）")
        preview_layout = QVBoxLayout()
        
        self.preview_table = QTableWidget()
        self.preview_table.setMaximumHeight(150)
        preview_layout.addWidget(self.preview_table)
        
        preview_group.setLayout(preview_layout)
        layout.addWidget(preview_group)
        
        # 功能说明区域
        info_group = QGroupBox("功能说明")
        info_layout = QVBoxLayout()
        
        info_label = QLabel("""
<b>Excel导入客户表说明:</b><br>
• 目标表: ba_rlb_customer<br>
• <b>两步导入流程:</b><br>
&nbsp;&nbsp;1. 先点「导入非企业法人」: 导入客户类型≠"企业法人"的记录，is_customer=1<br>
&nbsp;&nbsp;2. 再点「导入企业法人」: 导入客户类型="企业法人"的记录，如果不存在则is_customer=2<br>
• 法人姓名列: 必选，对应 legal_name 字段<br>
• 客户类型列: 必选，用于区分企业法人/非企业法人<br>
• 担当人员列: 可选，在 ba_admin 表中匹配 username 或 nickname<br>
• 导入前检查: 根据选择的字段检查是否已存在，存在则跳过<br>
• 自动填充: create_time, update_time, remark="系统导入"
        """)
        info_label.setWordWrap(True)
        info_layout.addWidget(info_label)
        
        info_group.setLayout(info_layout)
        layout.addWidget(info_group)
        
        # 操作按钮区域
        button_layout = QHBoxLayout()
        
        self.import_non_corporate_btn = QPushButton("第一步：导入非企业法人")
        self.import_non_corporate_btn.clicked.connect(self.start_import_non_corporate)
        self.import_non_corporate_btn.setEnabled(False)
        self.import_non_corporate_btn.setStyleSheet("background-color: #4CAF50; color: white; font-weight: bold; padding: 8px;")
        button_layout.addWidget(self.import_non_corporate_btn)
        
        self.import_corporate_btn = QPushButton("第二步：导入企业法人")
        self.import_corporate_btn.clicked.connect(self.start_import_corporate)
        self.import_corporate_btn.setEnabled(False)
        self.import_corporate_btn.setStyleSheet("background-color: #2196F3; color: white; font-weight: bold; padding: 8px;")
        button_layout.addWidget(self.import_corporate_btn)
        
        button_layout.addStretch()
        
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        button_layout.addWidget(self.progress_bar)
        
        layout.addLayout(button_layout)
        
        # 日志区域
        log_group = QGroupBox("导入日志")
        log_layout = QVBoxLayout()
        
        self.log_text = QTextEdit()
        self.log_text.setMaximumHeight(150)
        self.log_text.setReadOnly(True)
        log_layout.addWidget(self.log_text)
        
        log_button_layout = QHBoxLayout()
        self.clear_log_btn = QPushButton("清空日志")
        self.clear_log_btn.clicked.connect(self.clear_log)
        log_button_layout.addWidget(self.clear_log_btn)
        
        self.export_log_btn = QPushButton("导出日志")
        self.export_log_btn.clicked.connect(self.export_log)
        log_button_layout.addWidget(self.export_log_btn)
        
        log_button_layout.addStretch()
        log_layout.addLayout(log_button_layout)
        
        log_group.setLayout(log_layout)
        layout.addWidget(log_group)
        
        # 结果显示区域
        self.result_label = QLabel("")
        self.result_label.setAlignment(Qt.AlignCenter)
        font = QFont()
        font.setBold(True)
        self.result_label.setFont(font)
        layout.addWidget(self.result_label)
        
        self.setLayout(layout)
    
    def load_datasources(self):
        """加载数据源列表"""
        self.datasource_combo.clear()
        self.datasource_combo.addItem("请选择数据源", None)
        
        datasources = self.db_manager.get_all_datasources()
        for ds in datasources:
            self.datasource_combo.addItem(f"{ds.name} ({ds.host}:{ds.port})", ds)
    
    def on_datasource_changed(self):
        """数据源选择变化"""
        current_data = self.datasource_combo.currentData()
        self.test_connection_btn.setEnabled(current_data is not None)
        self.update_import_button_state()
    
    def test_connection(self):
        """测试数据库连接"""
        datasource = self.datasource_combo.currentData()
        if not datasource:
            return
        
        success, message = self.db_manager.test_connection(datasource)
        if success:
            QMessageBox.information(self, "连接成功", message)
        else:
            QMessageBox.critical(self, "连接失败", message)
    
    def select_input_file(self):
        """选择输入Excel文件"""
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "选择Excel文件",
            "",
            "Excel文件 (*.xlsx *.xls);;所有文件 (*)"
        )
        
        if file_path:
            self.input_file = file_path
            self.input_file_label.setText(os.path.basename(file_path))
            self.input_file_label.setStyleSheet("color: black; font-style: normal;")
            self.load_excel_preview()
            self.update_import_button_state()
    
    def load_excel_preview(self):
        """加载Excel文件预览"""
        try:
            if self.input_file.endswith('.xlsx'):
                try:
                    self.df = pd.read_excel(self.input_file, engine='openpyxl', header=0)
                except:
                    from openpyxl import load_workbook
                    wb = load_workbook(self.input_file, data_only=True)
                    ws = wb.active
                    
                    data = []
                    for row in ws.iter_rows(values_only=True):
                        data.append(row)
                    
                    if data:
                        columns = [str(col) if col is not None else f'列{i+1}' for i, col in enumerate(data[0])]
                        self.df = pd.DataFrame(data[1:], columns=columns)
            else:
                self.df = pd.read_excel(self.input_file, engine='xlrd', header=0)
            
            # 填充列选择下拉框
            self.populate_column_combos()
            
            # 显示数据预览
            self.show_data_preview()
            
        except Exception as e:
            QMessageBox.critical(self, "错误", f"读取Excel文件失败：{str(e)}")
    
    def populate_column_combos(self):
        """填充列选择下拉框"""
        if self.df is None:
            return
        
        columns = self.df.columns.tolist()
        
        col_options = []
        for i, col in enumerate(columns):
            col_letter = chr(65 + i) if i < 26 else f"A{chr(65 + i - 26)}" if i < 52 else f"Col{i+1}"
            col_options.append((f"{col_letter}列: {col}", i))
        
        # 法人姓名列（必选）
        self.legal_name_col_combo.clear()
        self.legal_name_col_combo.addItem("请选择列", -1)
        for text, idx in col_options:
            self.legal_name_col_combo.addItem(text, idx)
        self.legal_name_col_combo.setEnabled(True)
        
        # 客户类型列（必选）
        self.customer_type_col_combo.clear()
        self.customer_type_col_combo.addItem("请选择列", -1)
        for text, idx in col_options:
            self.customer_type_col_combo.addItem(text, idx)
        self.customer_type_col_combo.setEnabled(True)
        
        # 担当人员列（可选）
        self.staff_col_combo.clear()
        self.staff_col_combo.addItem("不使用", -1)
        for text, idx in col_options:
            self.staff_col_combo.addItem(text, idx)
        self.staff_col_combo.setEnabled(True)
        
        # 不自动匹配，让用户手动选择
        # self.auto_match_columns(columns)
    
    def auto_match_columns(self, columns):
        """尝试自动匹配常见列名"""
        for i, col in enumerate(columns):
            col_lower = str(col).lower()
            
            if '法人' in col_lower or 'legal' in col_lower:
                self.legal_name_col_combo.setCurrentIndex(i + 1)
            elif '类型' in col_lower or '属性' in col_lower:
                self.customer_type_col_combo.setCurrentIndex(i + 1)
            elif '担当' in col_lower or '负责' in col_lower or '人员' in col_lower:
                self.staff_col_combo.setCurrentIndex(i + 1)
    
    def show_data_preview(self):
        """显示数据预览"""
        preview_df = self.df.head(5)
        
        self.preview_table.setRowCount(len(preview_df))
        self.preview_table.setColumnCount(len(preview_df.columns))
        self.preview_table.setHorizontalHeaderLabels([str(col) for col in preview_df.columns])
        
        for i in range(len(preview_df)):
            for j, col in enumerate(preview_df.columns):
                value = preview_df.iloc[i, j]
                item = QTableWidgetItem(str(value) if pd.notna(value) else "")
                self.preview_table.setItem(i, j, item)
        
        self.preview_table.resizeColumnsToContents()
    
    def update_import_button_state(self):
        """更新导入按钮状态"""
        has_datasource = self.datasource_combo.currentData() is not None
        has_file = bool(self.input_file)
        legal_name_selected = (self.legal_name_col_combo.currentData() is not None and 
                               self.legal_name_col_combo.currentData() >= 0)
        customer_type_selected = (self.customer_type_col_combo.currentData() is not None and
                                  self.customer_type_col_combo.currentData() >= 0)
        
        can_import = has_datasource and has_file and legal_name_selected and customer_type_selected
        self.import_non_corporate_btn.setEnabled(can_import)
        self.import_corporate_btn.setEnabled(can_import)
    
    def get_column_mapping(self):
        """获取列映射配置"""
        return {
            'legal_name': self.legal_name_col_combo.currentData(),
            'customer_type': self.customer_type_col_combo.currentData() if self.customer_type_col_combo.currentData() >= 0 else None,
            'staff': self.staff_col_combo.currentData() if self.staff_col_combo.currentData() >= 0 else None
        }
    
    def start_import_non_corporate(self):
        """开始导入非企业法人"""
        self._start_import('non_corporate')
    
    def start_import_corporate(self):
        """开始导入企业法人"""
        self._start_import('corporate')
    
    def _start_import(self, import_mode):
        """开始导入"""
        datasource = self.datasource_combo.currentData()
        if not datasource or not self.input_file:
            QMessageBox.warning(self, "警告", "请选择数据源和Excel文件！")
            return
        
        column_mapping = self.get_column_mapping()
        check_field = self.check_field_combo.currentData()
        
        mode_name = "非企业法人" if import_mode == 'non_corporate' else "企业法人"
        is_customer_value = 1 if import_mode == 'non_corporate' else 2
        
        # 构建映射说明
        mapping_info = []
        mapping_info.append(f"法人姓名: {self.legal_name_col_combo.currentText()}")
        mapping_info.append(f"客户类型: {self.customer_type_col_combo.currentText()}")
        mapping_info.append(f"担当人员: {self.staff_col_combo.currentText()}")
        mapping_info.append(f"重复检查: {self.check_field_combo.currentText()}")
        mapping_info.append(f"导入模式: {mode_name}")
        mapping_info.append(f"is_customer值: {is_customer_value}")
        
        reply = QMessageBox.question(
            self,
            f"确认导入 - {mode_name}",
            f"确定要导入【{mode_name}】数据吗？\n\n"
            f"数据库: {datasource.name}\n"
            f"目标表: ba_rlb_customer\n\n"
            f"列映射配置:\n" + "\n".join(mapping_info) + "\n\n"
            f"• 只导入客户类型{'≠' if import_mode == 'non_corporate' else '='}「企业法人」的记录\n"
            f"• 新建的客户 is_customer={is_customer_value}\n"
            f"• 已存在的记录将被跳过",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )
        
        if reply != QMessageBox.Yes:
            return
        
        self.import_non_corporate_btn.setEnabled(False)
        self.import_corporate_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)
        self.result_label.setText(f"正在导入{mode_name}...")
        self.log_text.clear()
        
        self.worker = ExcelToDbWorker(self.input_file, datasource, self.db_manager, column_mapping, check_field, import_mode)
        self.worker.progress.connect(self.progress_bar.setValue)
        self.worker.log_message.connect(self.append_log)
        self.worker.finished.connect(self.on_import_finished)
        self.worker.error.connect(self.on_import_error)
        self.worker.start()
    
    def append_log(self, message):
        """添加日志消息"""
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.append(f"[{timestamp}] {message}")
    
    def on_import_finished(self, results):
        """导入完成"""
        self.update_import_button_state()
        self.progress_bar.setVisible(False)
        
        result_text = (
            f"导入完成！\n"
            f"总计：{results['total']} 行\n"
            f"成功：{results['success']} 行\n"
            f"跳过(类型不匹配)：{results['skipped_wrong_type']} 行\n"
            f"跳过(空值)：{results['skipped_empty']} 行\n"
            f"跳过(已存在)：{results['skipped_exists']} 行\n"
            f"失败：{results['failed']} 行"
        )
        
        self.result_label.setText(result_text)
        self.result_label.setStyleSheet("color: green;")
        
        QMessageBox.information(self, "导入完成", result_text)
    
    def on_import_error(self, error_msg):
        """导入错误"""
        self.update_import_button_state()
        self.progress_bar.setVisible(False)
        self.result_label.setText("导入失败！")
        self.result_label.setStyleSheet("color: red;")
        
        QMessageBox.critical(self, "导入错误", f"导入失败：{error_msg}")
    
    def clear_log(self):
        """清空日志"""
        self.log_text.clear()
    
    def export_log(self):
        """导出日志"""
        if not self.log_text.toPlainText():
            QMessageBox.information(self, "提示", "没有日志可以导出")
            return
        
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "保存日志文件",
            f"导入日志_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt",
            "文本文件 (*.txt);;所有文件 (*)"
        )
        
        if file_path:
            try:
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(self.log_text.toPlainText())
                QMessageBox.information(self, "成功", "日志导出成功！")
            except Exception as e:
                QMessageBox.critical(self, "错误", f"导出失败：{str(e)}")
