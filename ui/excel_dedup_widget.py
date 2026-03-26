"""
Excel表格去重功能界面
"""
import os
import pandas as pd
from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QFormLayout,
                             QPushButton, QLabel, QFileDialog, QMessageBox,
                             QTableWidget, QTableWidgetItem, QCheckBox,
                             QGroupBox, QProgressBar, QLineEdit, QScrollArea,
                             QFrame, QGridLayout)
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtGui import QFont


class ExcelDedupWorker(QThread):
    """Excel去重处理工作线程"""
    progress = pyqtSignal(int)
    finished = pyqtSignal(str, str, int, int, int)  # 去重文件路径, 重复文件路径, 原始行数, 去重后行数, 重复行数
    error = pyqtSignal(str)
    
    def __init__(self, input_file, output_file, duplicate_file, selected_columns):
        super().__init__()
        self.input_file = input_file
        self.output_file = output_file
        self.duplicate_file = duplicate_file
        self.selected_columns = selected_columns
    
    def run(self):
        try:
            self.progress.emit(10)
            
            # 读取Excel文件，确保读取所有列
            if self.input_file.endswith('.xlsx'):
                try:
                    df = pd.read_excel(self.input_file, engine='openpyxl', header=0)
                    
                    # 如果列数太少，手动读取
                    if len(df.columns) < 5:
                        from openpyxl import load_workbook
                        wb = load_workbook(self.input_file, data_only=True)
                        ws = wb.active
                        
                        max_row = ws.max_row
                        max_col = ws.max_column
                        
                        data = []
                        for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
                            data.append(row)
                        
                        if data:
                            columns = [str(col) if col is not None else f'列{i+1}' for i, col in enumerate(data[0])]
                            df = pd.DataFrame(data[1:], columns=columns)
                            
                except:
                    # 备用方案：转换为CSV
                    import tempfile
                    import os
                    temp_csv = tempfile.mktemp(suffix='.csv')
                    
                    from openpyxl import load_workbook
                    wb = load_workbook(self.input_file, data_only=True)
                    ws = wb.active
                    
                    import csv
                    with open(temp_csv, 'w', newline='', encoding='utf-8') as f:
                        writer = csv.writer(f)
                        for row in ws.iter_rows(values_only=True):
                            writer.writerow(row)
                    
                    df = pd.read_csv(temp_csv)
                    os.unlink(temp_csv)
            else:
                df = pd.read_excel(self.input_file, engine='xlrd', header=0)
            
            # 清理数据
            df = df.dropna(axis=1, how='all')  # 删除全空的列
            
            # 重命名未命名的列
            new_columns = []
            for i, col in enumerate(df.columns):
                if pd.isna(col) or str(col).startswith('Unnamed'):
                    new_columns.append(f'列{i+1}')
                else:
                    new_columns.append(str(col))
            df.columns = new_columns
            original_count = len(df)
            self.progress.emit(30)
            
            # 根据选择的列进行去重，并识别重复数据
            # 使用keep='last'保留最后一条（最新的）
            if self.selected_columns:
                # 标记重复行（keep='last'表示保留最后一条，前面的标记为重复）
                df['is_duplicate'] = df.duplicated(subset=self.selected_columns, keep='last')
                # 获取去重后的数据（保留最后一次出现的）
                df_dedup = df[~df['is_duplicate']].drop('is_duplicate', axis=1)
                # 获取重复的数据（除了最后一次出现的都是重复）
                df_duplicates = df[df['is_duplicate']].drop('is_duplicate', axis=1)
            else:
                # 对所有列去重
                df['is_duplicate'] = df.duplicated(keep='last')
                df_dedup = df[~df['is_duplicate']].drop('is_duplicate', axis=1)
                df_duplicates = df[df['is_duplicate']].drop('is_duplicate', axis=1)
            
            dedup_count = len(df_dedup)
            duplicate_count = len(df_duplicates)
            self.progress.emit(50)
            
            # 导出去重后的数据
            df_dedup.to_excel(self.output_file, index=False)
            self.progress.emit(75)
            
            # 导出重复的数据
            if duplicate_count > 0:
                # 为重复数据添加额外信息
                df_duplicates_with_info = df_duplicates.copy()
                
                # 添加重复原因列
                if self.selected_columns:
                    duplicate_reason = f"基于列 [{', '.join(self.selected_columns)}] 的重复数据（保留了最后一条）"
                else:
                    duplicate_reason = "基于所有列的重复数据（保留了最后一条）"
                
                df_duplicates_with_info.insert(0, '重复原因', duplicate_reason)
                df_duplicates_with_info.insert(1, '原始行号', df_duplicates.index + 2)  # +2因为Excel从1开始且有表头
                
                # 导出重复数据
                df_duplicates_with_info.to_excel(self.duplicate_file, index=False)
            else:
                # 如果没有重复数据，创建空文件
                empty_df = pd.DataFrame(columns=['说明'])
                empty_df.loc[0] = ['没有发现重复数据']
                empty_df.to_excel(self.duplicate_file, index=False)
            
            self.progress.emit(100)
            
            self.finished.emit(self.output_file, self.duplicate_file, original_count, dedup_count, duplicate_count)
            
        except Exception as e:
            self.error.emit(str(e))


class ExcelDedupWidget(QWidget):
    """Excel表格去重功能组件"""
    
    def __init__(self):
        super().__init__()
        self.input_file = ""
        self.output_file = ""
        self.df = None
        self.column_checkboxes = {}
        self.init_ui()
    
    def init_ui(self):
        """初始化UI"""
        layout = QVBoxLayout()
        
        # 文件选择区域
        file_group = QGroupBox("文件选择")
        file_layout = QFormLayout()
        
        # 输入文件
        input_layout = QHBoxLayout()
        self.input_file_label = QLabel("请选择Excel文件...")
        self.input_file_label.setStyleSheet("color: gray; font-style: italic;")
        input_layout.addWidget(self.input_file_label)
        
        self.select_input_btn = QPushButton("选择输入文件")
        self.select_input_btn.clicked.connect(self.select_input_file)
        input_layout.addWidget(self.select_input_btn)
        
        file_layout.addRow("输入文件:", input_layout)
        
        # 去重后输出文件
        output_layout = QHBoxLayout()
        self.output_file_edit = QLineEdit()
        self.output_file_edit.setPlaceholderText("去重后文件路径（可选，默认在输入文件同目录）")
        output_layout.addWidget(self.output_file_edit)
        
        self.select_output_btn = QPushButton("选择输出位置")
        self.select_output_btn.clicked.connect(self.select_output_file)
        output_layout.addWidget(self.select_output_btn)
        
        file_layout.addRow("去重后文件:", output_layout)
        
        # 重复数据输出文件
        duplicate_layout = QHBoxLayout()
        self.duplicate_file_edit = QLineEdit()
        self.duplicate_file_edit.setPlaceholderText("重复数据文件路径（可选，默认在输入文件同目录）")
        duplicate_layout.addWidget(self.duplicate_file_edit)
        
        self.select_duplicate_btn = QPushButton("选择输出位置")
        self.select_duplicate_btn.clicked.connect(self.select_duplicate_file)
        duplicate_layout.addWidget(self.select_duplicate_btn)
        
        file_layout.addRow("重复数据文件:", duplicate_layout)
        
        file_group.setLayout(file_layout)
        layout.addWidget(file_group)
        
        # 列选择区域
        self.column_group = QGroupBox("选择去重依据列（不选择则对所有列去重，保留最后一条/最新的）")
        self.column_group.setEnabled(False)
        
        # 使用滚动区域来容纳可能很多的列
        scroll_area = QScrollArea()
        scroll_widget = QWidget()
        self.column_layout = QGridLayout(scroll_widget)
        scroll_area.setWidget(scroll_widget)
        scroll_area.setWidgetResizable(True)
        scroll_area.setMaximumHeight(200)
        
        column_group_layout = QVBoxLayout()
        
        # 全选/取消全选按钮
        select_all_layout = QHBoxLayout()
        self.select_all_btn = QPushButton("全选")
        self.select_all_btn.clicked.connect(self.select_all_columns)
        select_all_layout.addWidget(self.select_all_btn)
        
        self.deselect_all_btn = QPushButton("取消全选")
        self.deselect_all_btn.clicked.connect(self.deselect_all_columns)
        select_all_layout.addWidget(self.deselect_all_btn)
        
        select_all_layout.addStretch()
        column_group_layout.addLayout(select_all_layout)
        
        column_group_layout.addWidget(scroll_area)
        self.column_group.setLayout(column_group_layout)
        layout.addWidget(self.column_group)
        
        # 数据预览区域
        self.preview_group = QGroupBox("数据预览（前10行）")
        self.preview_group.setEnabled(False)
        preview_layout = QVBoxLayout()
        
        self.preview_table = QTableWidget()
        self.preview_table.setMaximumHeight(300)
        preview_layout.addWidget(self.preview_table)
        
        self.preview_group.setLayout(preview_layout)
        layout.addWidget(self.preview_group)
        
        # 操作按钮区域
        button_layout = QHBoxLayout()
        
        self.process_btn = QPushButton("开始去重处理")
        self.process_btn.clicked.connect(self.start_dedup_process)
        self.process_btn.setEnabled(False)
        button_layout.addWidget(self.process_btn)
        
        button_layout.addStretch()
        
        # 进度条
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        button_layout.addWidget(self.progress_bar)
        
        layout.addLayout(button_layout)
        
        # 结果显示区域
        self.result_label = QLabel("")
        self.result_label.setAlignment(Qt.AlignCenter)
        font = QFont()
        font.setBold(True)
        self.result_label.setFont(font)
        layout.addWidget(self.result_label)
        
        layout.addStretch()
        self.setLayout(layout)
    
    def select_input_file(self):
        """选择输入Excel文件"""
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "选择Excel文件",
            "",
            "Excel文件 (*.xlsx *.xls);;所有文件 (*)"
        )
        
        if file_path:
            self.input_file = file_path
            self.input_file_label.setText(os.path.basename(file_path))
            self.input_file_label.setStyleSheet("color: black; font-style: normal;")
            
            # 自动设置输出文件路径
            if not self.output_file_edit.text():
                base_name = os.path.splitext(file_path)[0]
                self.output_file_edit.setText(f"{base_name}_去重.xlsx")
            
            # 自动设置重复数据文件路径
            if not self.duplicate_file_edit.text():
                base_name = os.path.splitext(file_path)[0]
                self.duplicate_file_edit.setText(f"{base_name}_重复数据.xlsx")
            
            self.load_excel_preview()
    
    def select_output_file(self):
        """选择输出文件位置"""
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "选择输出文件位置",
            self.output_file_edit.text() or "去重结果.xlsx",
            "Excel文件 (*.xlsx);;所有文件 (*)"
        )
        
        if file_path:
            self.output_file_edit.setText(file_path)
    
    def select_duplicate_file(self):
        """选择重复数据文件位置"""
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "选择重复数据文件位置",
            self.duplicate_file_edit.text() or "重复数据.xlsx",
            "Excel文件 (*.xlsx);;所有文件 (*)"
        )
        
        if file_path:
            self.duplicate_file_edit.setText(file_path)
    
    def load_excel_preview(self):
        """加载Excel文件预览"""
        try:
            # 读取Excel文件，确保读取所有列
            if self.input_file.endswith('.xlsx'):
                # 对于xlsx文件，尝试多种方式
                try:
                    # 先尝试openpyxl，不限制列数
                    self.df = pd.read_excel(self.input_file, engine='openpyxl', header=0)
                    
                    # 如果列数太少，可能是读取有问题，尝试手动读取
                    if len(self.df.columns) < 5:  # 如果列数少于5，可能有问题
                        from openpyxl import load_workbook
                        wb = load_workbook(self.input_file, data_only=True)
                        ws = wb.active
                        
                        # 获取实际的数据范围
                        max_row = ws.max_row
                        max_col = ws.max_column
                        
                        print(f"Excel实际范围: {max_row}行 x {max_col}列")
                        
                        # 手动读取所有数据
                        data = []
                        for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
                            data.append(row)
                        
                        if data:
                            # 第一行作为列名
                            columns = [str(col) if col is not None else f'列{i+1}' for i, col in enumerate(data[0])]
                            # 剩余行作为数据
                            self.df = pd.DataFrame(data[1:], columns=columns)
                        
                except Exception as e:
                    print(f"openpyxl读取失败: {e}")
                    # 备用方案：转换为csv再读取
                    import tempfile
                    import os
                    temp_csv = tempfile.mktemp(suffix='.csv')
                    
                    from openpyxl import load_workbook
                    wb = load_workbook(self.input_file, data_only=True)
                    ws = wb.active
                    
                    import csv
                    with open(temp_csv, 'w', newline='', encoding='utf-8') as f:
                        writer = csv.writer(f)
                        for row in ws.iter_rows(values_only=True):
                            writer.writerow(row)
                    
                    self.df = pd.read_csv(temp_csv)
                    os.unlink(temp_csv)
            else:
                # 对于xls文件，使用xlrd
                self.df = pd.read_excel(self.input_file, engine='xlrd', header=0)
            
            # 清理列名，移除空列
            self.df = self.df.dropna(axis=1, how='all')  # 删除全空的列
            
            # 重命名未命名的列
            new_columns = []
            for i, col in enumerate(self.df.columns):
                if pd.isna(col) or str(col).startswith('Unnamed'):
                    new_columns.append(f'列{i+1}')
                else:
                    new_columns.append(str(col))
            self.df.columns = new_columns
            
            print(f"最终读取到 {len(self.df)} 行 x {len(self.df.columns)} 列")
            print(f"列名: {list(self.df.columns)}")
            
            # 启用相关组件
            self.column_group.setEnabled(True)
            self.preview_group.setEnabled(True)
            self.process_btn.setEnabled(True)
            
            # 创建列选择复选框
            self.create_column_checkboxes()
            
            # 显示数据预览
            self.show_data_preview()
            
        except Exception as e:
            QMessageBox.critical(self, "错误", f"读取Excel文件失败：{str(e)}")
    
    def create_column_checkboxes(self):
        """创建列选择复选框"""
        # 清除现有的复选框
        for checkbox in self.column_checkboxes.values():
            checkbox.deleteLater()
        self.column_checkboxes.clear()
        
        # 清除布局中的所有项目
        for i in reversed(range(self.column_layout.count())):
            self.column_layout.itemAt(i).widget().setParent(None)
        
        # 创建新的复选框
        columns = self.df.columns.tolist()
        cols_per_row = 3  # 每行显示3个复选框
        
        for i, column in enumerate(columns):
            checkbox = QCheckBox(str(column))
            checkbox.setChecked(True)  # 默认全选
            self.column_checkboxes[column] = checkbox
            
            row = i // cols_per_row
            col = i % cols_per_row
            self.column_layout.addWidget(checkbox, row, col)
    
    def show_data_preview(self):
        """显示数据预览"""
        # 显示前10行数据
        preview_df = self.df.head(10)
        
        self.preview_table.setRowCount(len(preview_df))
        self.preview_table.setColumnCount(len(preview_df.columns))
        self.preview_table.setHorizontalHeaderLabels([str(col) for col in preview_df.columns])
        
        for i in range(len(preview_df)):
            for j, col in enumerate(preview_df.columns):
                item = QTableWidgetItem(str(preview_df.iloc[i, j]))
                self.preview_table.setItem(i, j, item)
        
        # 自适应列宽
        self.preview_table.resizeColumnsToContents()
    
    def select_all_columns(self):
        """全选所有列"""
        for checkbox in self.column_checkboxes.values():
            checkbox.setChecked(True)
    
    def deselect_all_columns(self):
        """取消全选所有列"""
        for checkbox in self.column_checkboxes.values():
            checkbox.setChecked(False)
    
    def get_selected_columns(self):
        """获取选中的列"""
        selected = []
        for column, checkbox in self.column_checkboxes.items():
            if checkbox.isChecked():
                selected.append(column)
        return selected
    
    def start_dedup_process(self):
        """开始去重处理"""
        if not self.input_file:
            QMessageBox.warning(self, "警告", "请先选择输入文件！")
            return
        
        output_file = self.output_file_edit.text().strip()
        if not output_file:
            QMessageBox.warning(self, "警告", "请设置去重后文件路径！")
            return
        
        duplicate_file = self.duplicate_file_edit.text().strip()
        if not duplicate_file:
            QMessageBox.warning(self, "警告", "请设置重复数据文件路径！")
            return
        
        selected_columns = self.get_selected_columns()
        
        # 确认对话框
        if selected_columns:
            msg = f"将根据以下列进行去重：\n{', '.join(selected_columns)}\n\n"
        else:
            msg = "将对所有列进行去重\n\n"
        
        msg += f"去重后数据将保存到：{os.path.basename(output_file)}\n"
        msg += f"重复数据将保存到：{os.path.basename(duplicate_file)}\n\n"
        msg += "确定开始处理吗？"
        
        reply = QMessageBox.question(self, "确认", msg, QMessageBox.Yes | QMessageBox.No)
        if reply != QMessageBox.Yes:
            return
        
        # 禁用按钮，显示进度条
        self.process_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)
        self.result_label.setText("正在处理中...")
        
        # 启动工作线程
        self.worker = ExcelDedupWorker(self.input_file, output_file, duplicate_file, selected_columns)
        self.worker.progress.connect(self.progress_bar.setValue)
        self.worker.finished.connect(self.on_process_finished)
        self.worker.error.connect(self.on_process_error)
        self.worker.start()
    
    def on_process_finished(self, output_file, duplicate_file, original_count, dedup_count, duplicate_count):
        """处理完成"""
        self.process_btn.setEnabled(True)
        self.progress_bar.setVisible(False)
        
        result_text = f"处理完成！\n原始数据：{original_count} 行\n去重后：{dedup_count} 行\n重复数据：{duplicate_count} 行"
        self.result_label.setText(result_text)
        self.result_label.setStyleSheet("color: green;")
        
        # 详细的完成信息
        detail_msg = f"{result_text}\n\n"
        detail_msg += f"去重后文件：{os.path.basename(output_file)}\n"
        detail_msg += f"重复数据文件：{os.path.basename(duplicate_file)}\n\n"
        
        if duplicate_count > 0:
            detail_msg += f"发现 {duplicate_count} 行重复数据，已保存到重复数据文件中供您分析对比。\n\n"
        else:
            detail_msg += "没有发现重复数据。\n\n"
        
        detail_msg += "是否打开输出文件所在文件夹？"
        
        # 询问是否打开输出文件夹
        reply = QMessageBox.question(
            self,
            "处理完成",
            detail_msg,
            QMessageBox.Yes | QMessageBox.No
        )
        
        if reply == QMessageBox.Yes:
            import subprocess
            subprocess.Popen(f'explorer /select,"{output_file}"')
    
    def on_process_error(self, error_msg):
        """处理错误"""
        self.process_btn.setEnabled(True)
        self.progress_bar.setVisible(False)
        self.result_label.setText("处理失败！")
        self.result_label.setStyleSheet("color: red;")
        
        QMessageBox.critical(self, "处理错误", f"去重处理失败：{error_msg}")
