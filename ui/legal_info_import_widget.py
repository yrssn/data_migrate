"""
法人信息导入功能界面
"""
import os
import pandas as pd
from datetime import datetime
from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QFormLayout,
                             QPushButton, QLabel, QFileDialog, QMessageBox,
                             QTableWidget, QTableWidgetItem, QComboBox,
                             QGroupBox, QProgressBar, QLineEdit, QTextEdit,
                             QCheckBox, QSpinBox)
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtGui import QFont
from database import DatabaseManager
import pymysql


class LegalInfoImportWorker(QThread):
    """法人信息导入工作线程"""
    progress = pyqtSignal(int)
    log_message = pyqtSignal(str)
    finished = pyqtSignal(dict)  # 返回统计结果
    error = pyqtSignal(str)
    
    def __init__(self, input_file, datasource, db_manager):
        super().__init__()
        self.input_file = input_file
        self.datasource = datasource
        self.db_manager = db_manager
        self.results = {
            'total': 0,
            'success_jp': 0,
            'success_non_jp': 0,
            'skipped_not_corporate': 0,
            'skipped_customer_not_found': 0,
            'skipped_duplicate': 0,  # 新增：跳过重复记录
            'failed': 0,
            'error_records': []  # 存储错误记录用于导出
        }
    
    def run(self):
        try:
            self.log_message.emit("开始读取Excel文件...")
            self.progress.emit(5)
            
            # 读取Excel文件
            if self.input_file.endswith('.xlsx'):
                try:
                    df = pd.read_excel(self.input_file, engine='openpyxl', header=0)
                except:
                    from openpyxl import load_workbook
                    wb = load_workbook(self.input_file, data_only=True)
                    ws = wb.active
                    
                    data = []
                    for row in ws.iter_rows(values_only=True):
                        data.append(row)
                    
                    if data:
                        columns = [str(col) if col is not None else f'列{i+1}' for i, col in enumerate(data[0])]
                        df = pd.DataFrame(data[1:], columns=columns)
            else:
                df = pd.read_excel(self.input_file, engine='xlrd', header=0)
            
            self.results['total'] = len(df)
            self.log_message.emit(f"读取到 {len(df)} 行数据")
            self.progress.emit(10)
            
            # 连接数据库
            self.log_message.emit("连接数据库...")
            connection = pymysql.connect(
                host=self.datasource.host,
                port=self.datasource.port,
                user=self.datasource.username,
                password=self.datasource.password,
                database=self.datasource.database,
                charset=self.datasource.charset
            )
            
            cursor = connection.cursor()
            self.progress.emit(15)
            
            # 处理每一行数据
            for index, row in df.iterrows():
                try:
                    progress = 15 + int((index / len(df)) * 80)
                    self.progress.emit(progress)
                    
                    # 检查列数是否足够
                    if len(df.columns) < 8:
                        error_msg = f"第{index+2}行: Excel列数不足，需要至少8列"
                        self.log_message.emit(error_msg)
                        self.add_error_record(df, index, error_msg)
                        self.results['failed'] += 1
                        continue
                    
                    # 获取G列数据（客户类型）
                    g_value = str(row.iloc[6]).strip() if pd.notna(row.iloc[6]) else ""
                    if g_value != "企业法人":
                        error_msg = f"第{index+2}行: G列不是'企业法人'，跳过 (值: '{g_value}')"
                        self.log_message.emit(error_msg)
                        self.add_error_record(df, index, error_msg)
                        self.results['skipped_not_corporate'] += 1
                        continue
                    
                    # 获取B列数据（法人姓名）
                    legal_name = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ""
                    if not legal_name:
                        error_msg = f"第{index+2}行: B列法人姓名为空"
                        self.log_message.emit(error_msg)
                        self.add_error_record(df, index, error_msg)
                        self.results['failed'] += 1
                        continue
                    
                    # 在ba_rlb_customer表中查找对应的客户
                    cursor.execute("""
                        SELECT id, admin_id, rlb_staff_id, admin_dept_id 
                        FROM ba_rlb_customer 
                        WHERE legal_name = %s
                    """, (legal_name,))
                    customer_result = cursor.fetchone()
                    
                    if not customer_result:
                        error_msg = f"第{index+2}行: 未找到法人姓名 '{legal_name}' 对应的客户记录"
                        self.log_message.emit(error_msg)
                        self.add_error_record(df, index, error_msg)
                        self.results['skipped_customer_not_found'] += 1
                        continue
                    
                    customer_id, admin_id, rlb_staff_id, admin_dept_id = customer_result
                    self.log_message.emit(f"第{index+2}行: 找到客户 '{legal_name}' (ID: {customer_id})")
                    
                    # 获取F列数据（国家）
                    country = str(row.iloc[5]).strip() if pd.notna(row.iloc[5]) else ""
                    
                    # 获取E列数据（公司名）
                    company_name = str(row.iloc[4]).strip() if pd.notna(row.iloc[4]) else ""
                    
                    current_time = int(datetime.now().timestamp())
                    
                    # 根据F列值（国家）决定插入到哪个表
                    if country.upper() == "日本" or country.upper() == "JAPAN" or country.upper() == "JP":
                        # 检查ba_rlb_legal_information表中是否已存在相同记录
                        cursor.execute("""
                            SELECT id FROM ba_rlb_legal_information 
                            WHERE company_name = %s AND country = %s AND legal_id = %s
                        """, (company_name, country, customer_id))
                        existing_record = cursor.fetchone()
                        
                        if existing_record:
                            error_msg = f"第{index+2}行: ba_rlb_legal_information表中已存在相同记录 - 公司: '{company_name}', 国家: '{country}', 法人ID: {customer_id}"
                            self.log_message.emit(error_msg)
                            self.add_error_record(df, index, error_msg)
                            self.results['skipped_duplicate'] += 1
                            continue
                        
                        # 插入到ba_rlb_legal_information表
                        insert_sql = """
                            INSERT INTO ba_rlb_legal_information 
                            (company_name, country, legal_id, admin_id, create_staff, rlb_staff_id, create_time, update_time)
                            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                        """
                        
                        cursor.execute(insert_sql, (
                            company_name,
                            country,
                            customer_id,
                            admin_id,
                            "系统导入",  # create_staff使用固定值"系统导入"
                            rlb_staff_id,
                            current_time,
                            current_time
                        ))
                        
                        self.results['success_jp'] += 1
                        self.log_message.emit(f"第{index+2}行: 成功插入到ba_rlb_legal_information表 - 公司: '{company_name}', 国家: '{country}'")
                        
                    else:
                        # 检查ba_rlb_legal_information_part2表中是否已存在相同记录
                        cursor.execute("""
                            SELECT id FROM ba_rlb_legal_information_part2 
                            WHERE company_name = %s AND country = %s AND legal_id = %s
                        """, (company_name, country, customer_id))
                        existing_record = cursor.fetchone()
                        
                        if existing_record:
                            error_msg = f"第{index+2}行: ba_rlb_legal_information_part2表中已存在相同记录 - 公司: '{company_name}', 国家: '{country}', 法人ID: {customer_id}"
                            self.log_message.emit(error_msg)
                            self.add_error_record(df, index, error_msg)
                            self.results['skipped_duplicate'] += 1
                            continue
                        
                        # 插入到ba_rlb_legal_information_part2表
                        insert_sql = """
                            INSERT INTO ba_rlb_legal_information_part2 
                            (company_name, country, legal_id, admin_id, create_staff, rlb_staff_id, create_time, update_time)
                            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                        """
                        
                        cursor.execute(insert_sql, (
                            company_name,
                            country,
                            customer_id,
                            admin_id,
                            "系统导入",  # create_staff使用固定值"系统导入"
                            rlb_staff_id,
                            current_time,
                            current_time
                        ))
                        
                        self.results['success_non_jp'] += 1
                        self.log_message.emit(f"第{index+2}行: 成功插入到ba_rlb_legal_information_part2表 - 公司: '{company_name}', 国家: '{country}'")
                    
                except Exception as e:
                    self.results['failed'] += 1
                    error_msg = f"第{index+2}行: 处理失败 - {str(e)}"
                    self.log_message.emit(error_msg)
                    self.add_error_record(df, index, error_msg)
            
            # 提交事务
            connection.commit()
            cursor.close()
            connection.close()
            
            self.progress.emit(100)
            self.finished.emit(self.results)
            
        except Exception as e:
            self.error.emit(str(e))
    
    def add_error_record(self, df, index, error_msg):
        """添加错误记录"""
        row_data = df.iloc[index].to_dict()
        row_data['行号'] = index + 2
        row_data['错误原因'] = error_msg
        self.results['error_records'].append(row_data)


class LegalInfoImportWidget(QWidget):
    """法人信息导入功能组件"""
    
    def __init__(self, db_manager: DatabaseManager):
        super().__init__()
        self.db_manager = db_manager
        self.input_file = ""
        self.df = None
        self.error_records = []
        self.init_ui()
        self.load_datasources()
    
    def init_ui(self):
        """初始化UI"""
        layout = QVBoxLayout()
        
        # 数据源选择区域
        datasource_group = QGroupBox("选择目标数据库")
        datasource_layout = QFormLayout()
        
        self.datasource_combo = QComboBox()
        self.datasource_combo.currentTextChanged.connect(self.on_datasource_changed)
        datasource_layout.addRow("数据源:", self.datasource_combo)
        
        self.test_connection_btn = QPushButton("测试连接")
        self.test_connection_btn.clicked.connect(self.test_connection)
        self.test_connection_btn.setEnabled(False)
        datasource_layout.addRow("", self.test_connection_btn)
        
        datasource_group.setLayout(datasource_layout)
        layout.addWidget(datasource_group)
        
        # 文件选择区域
        file_group = QGroupBox("选择Excel文件")
        file_layout = QFormLayout()
        
        input_layout = QHBoxLayout()
        self.input_file_label = QLabel("请选择Excel文件...")
        self.input_file_label.setStyleSheet("color: gray; font-style: italic;")
        input_layout.addWidget(self.input_file_label)
        
        self.select_file_btn = QPushButton("选择文件")
        self.select_file_btn.clicked.connect(self.select_input_file)
        input_layout.addWidget(self.select_file_btn)
        
        file_layout.addRow("Excel文件:", input_layout)
        file_group.setLayout(file_layout)
        layout.addWidget(file_group)
        
        # 数据预览区域
        self.preview_group = QGroupBox("数据预览和处理逻辑说明")
        self.preview_group.setEnabled(False)
        preview_layout = QVBoxLayout()
        
        # 说明文字
        info_label = QLabel("""
<b>处理逻辑说明:</b><br>
• <b>G列</b>: 必须是"企业法人"，否则跳过<br>
• <b>B列</b>: 法人姓名 → 在ba_rlb_customer表中查找对应ID<br>
• <b>F列</b>: 国家判断<br>
  &nbsp;&nbsp;- 如果是"日本"、"JAPAN"或"JP" → 插入ba_rlb_legal_information表<br>
  &nbsp;&nbsp;- 如果不是日本 → 插入ba_rlb_legal_information_part2表<br>
• <b>F列</b>: 国家 → country字段<br>
• <b>E列</b>: 公司名 → company_name字段<br><br>
<b>字段映射:</b><br>
• legal_id ← ba_rlb_customer.id<br>
• admin_id ← ba_rlb_customer.admin_id<br>
• rlb_staff_id ← ba_rlb_customer.rlb_staff_id<br>
• create_staff ← 固定值"系统导入"<br>
• create_time, update_time ← 当前时间戳
        """)
        info_label.setWordWrap(True)
        preview_layout.addWidget(info_label)
        
        self.preview_table = QTableWidget()
        self.preview_table.setMaximumHeight(200)
        preview_layout.addWidget(self.preview_table)
        
        self.preview_group.setLayout(preview_layout)
        layout.addWidget(self.preview_group)
        
        # 操作按钮区域
        button_layout = QHBoxLayout()
        
        self.import_btn = QPushButton("开始导入")
        self.import_btn.clicked.connect(self.start_import)
        self.import_btn.setEnabled(False)
        button_layout.addWidget(self.import_btn)
        
        button_layout.addStretch()
        
        # 进度条
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        button_layout.addWidget(self.progress_bar)
        
        layout.addLayout(button_layout)
        
        # 日志区域
        log_group = QGroupBox("导入日志")
        log_layout = QVBoxLayout()
        
        self.log_text = QTextEdit()
        self.log_text.setMaximumHeight(200)
        self.log_text.setReadOnly(True)
        log_layout.addWidget(self.log_text)
        
        log_button_layout = QHBoxLayout()
        self.clear_log_btn = QPushButton("清空日志")
        self.clear_log_btn.clicked.connect(self.clear_log)
        log_button_layout.addWidget(self.clear_log_btn)
        
        self.export_log_btn = QPushButton("导出日志")
        self.export_log_btn.clicked.connect(self.export_log)
        log_button_layout.addWidget(self.export_log_btn)
        
        self.export_errors_btn = QPushButton("导出错误记录Excel")
        self.export_errors_btn.clicked.connect(self.export_error_records)
        self.export_errors_btn.setEnabled(False)
        log_button_layout.addWidget(self.export_errors_btn)
        
        log_button_layout.addStretch()
        log_layout.addLayout(log_button_layout)
        
        log_group.setLayout(log_layout)
        layout.addWidget(log_group)
        
        # 结果显示区域
        self.result_label = QLabel("")
        self.result_label.setAlignment(Qt.AlignCenter)
        font = QFont()
        font.setBold(True)
        self.result_label.setFont(font)
        layout.addWidget(self.result_label)
        
        self.setLayout(layout)
    
    def load_datasources(self):
        """加载数据源列表"""
        self.datasource_combo.clear()
        self.datasource_combo.addItem("请选择数据源", None)
        
        datasources = self.db_manager.get_all_datasources()
        for ds in datasources:
            self.datasource_combo.addItem(f"{ds.name} ({ds.host}:{ds.port})", ds)
    
    def on_datasource_changed(self):
        """数据源选择变化"""
        current_data = self.datasource_combo.currentData()
        self.test_connection_btn.setEnabled(current_data is not None)
        self.update_import_button_state()
    
    def test_connection(self):
        """测试数据库连接"""
        datasource = self.datasource_combo.currentData()
        if not datasource:
            return
        
        success, message = self.db_manager.test_connection(datasource)
        if success:
            QMessageBox.information(self, "连接成功", message)
        else:
            QMessageBox.critical(self, "连接失败", message)
    
    def select_input_file(self):
        """选择输入Excel文件"""
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "选择Excel文件",
            "",
            "Excel文件 (*.xlsx *.xls);;所有文件 (*)"
        )
        
        if file_path:
            self.input_file = file_path
            self.input_file_label.setText(os.path.basename(file_path))
            self.input_file_label.setStyleSheet("color: black; font-style: normal;")
            self.load_excel_preview()
            self.update_import_button_state()
    
    def load_excel_preview(self):
        """加载Excel文件预览"""
        try:
            # 读取Excel文件
            if self.input_file.endswith('.xlsx'):
                try:
                    self.df = pd.read_excel(self.input_file, engine='openpyxl', header=0)
                except:
                    from openpyxl import load_workbook
                    wb = load_workbook(self.input_file, data_only=True)
                    ws = wb.active
                    
                    data = []
                    for row in ws.iter_rows(values_only=True):
                        data.append(row)
                    
                    if data:
                        columns = [str(col) if col is not None else f'列{i+1}' for i, col in enumerate(data[0])]
                        self.df = pd.DataFrame(data[1:], columns=columns)
            else:
                self.df = pd.read_excel(self.input_file, engine='xlrd', header=0)
            
            # 启用预览组件
            self.preview_group.setEnabled(True)
            
            # 显示数据预览（前5行）
            self.show_data_preview()
            
        except Exception as e:
            QMessageBox.critical(self, "错误", f"读取Excel文件失败：{str(e)}")
    
    def show_data_preview(self):
        """显示数据预览"""
        preview_df = self.df.head(5)
        
        self.preview_table.setRowCount(len(preview_df))
        self.preview_table.setColumnCount(len(preview_df.columns))
        self.preview_table.setHorizontalHeaderLabels([str(col) for col in preview_df.columns])
        
        for i in range(len(preview_df)):
            for j, col in enumerate(preview_df.columns):
                item = QTableWidgetItem(str(preview_df.iloc[i, j]))
                self.preview_table.setItem(i, j, item)
        
        # 自适应列宽
        self.preview_table.resizeColumnsToContents()
    
    def update_import_button_state(self):
        """更新导入按钮状态"""
        has_datasource = self.datasource_combo.currentData() is not None
        has_file = bool(self.input_file)
        self.import_btn.setEnabled(has_datasource and has_file)
    
    def start_import(self):
        """开始导入"""
        datasource = self.datasource_combo.currentData()
        if not datasource or not self.input_file:
            QMessageBox.warning(self, "警告", "请选择数据源和Excel文件！")
            return
        
        # 确认对话框
        reply = QMessageBox.question(
            self,
            "确认导入",
            f"确定要将Excel数据导入到数据库 '{datasource.name}' 吗？\n\n"
            f"数据将根据F列（国家）值插入到不同的法人信息表中：\n"
            f"• 日本/JAPAN/JP → ba_rlb_legal_information表\n"
            f"• 其他国家 → ba_rlb_legal_information_part2表",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )
        
        if reply != QMessageBox.Yes:
            return
        
        # 禁用按钮，显示进度条
        self.import_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setValue(0)
        self.result_label.setText("正在导入中...")
        self.log_text.clear()
        self.error_records = []
        self.export_errors_btn.setEnabled(False)
        
        # 启动工作线程
        self.worker = LegalInfoImportWorker(self.input_file, datasource, self.db_manager)
        self.worker.progress.connect(self.progress_bar.setValue)
        self.worker.log_message.connect(self.append_log)
        self.worker.finished.connect(self.on_import_finished)
        self.worker.error.connect(self.on_import_error)
        self.worker.start()
    
    def append_log(self, message):
        """添加日志消息"""
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.append(f"[{timestamp}] {message}")
    
    def on_import_finished(self, results):
        """导入完成"""
        self.import_btn.setEnabled(True)
        self.progress_bar.setVisible(False)
        self.error_records = results['error_records']
        
        result_text = (
            f"导入完成！\n"
            f"总计：{results['total']} 行\n"
            f"成功插入JP表：{results['success_jp']} 行\n"
            f"成功插入非JP表：{results['success_non_jp']} 行\n"
            f"跳过(非企业法人)：{results['skipped_not_corporate']} 行\n"
            f"跳过(客户未找到)：{results['skipped_customer_not_found']} 行\n"
            f"跳过(重复记录)：{results['skipped_duplicate']} 行\n"
            f"失败：{results['failed']} 行"
        )
        
        self.result_label.setText(result_text)
        self.result_label.setStyleSheet("color: green;")
        
        # 如果有错误记录，启用导出按钮
        if self.error_records:
            self.export_errors_btn.setEnabled(True)
        
        QMessageBox.information(self, "导入完成", result_text)
    
    def on_import_error(self, error_msg):
        """导入错误"""
        self.import_btn.setEnabled(True)
        self.progress_bar.setVisible(False)
        self.result_label.setText("导入失败！")
        self.result_label.setStyleSheet("color: red;")
        
        QMessageBox.critical(self, "导入错误", f"导入失败：{error_msg}")
    
    def clear_log(self):
        """清空日志"""
        self.log_text.clear()
    
    def export_log(self):
        """导出日志"""
        if not self.log_text.toPlainText():
            QMessageBox.information(self, "提示", "没有日志可以导出")
            return
        
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "保存日志文件",
            f"法人信息导入日志_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt",
            "文本文件 (*.txt);;所有文件 (*)"
        )
        
        if file_path:
            try:
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(self.log_text.toPlainText())
                QMessageBox.information(self, "成功", "日志导出成功！")
            except Exception as e:
                QMessageBox.critical(self, "错误", f"导出失败：{str(e)}")
    
    def export_error_records(self):
        """导出错误记录Excel"""
        if not self.error_records:
            QMessageBox.information(self, "提示", "没有错误记录可以导出")
            return
        
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "保存错误记录文件",
            f"法人信息导入错误记录_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            "Excel文件 (*.xlsx);;所有文件 (*)"
        )
        
        if file_path:
            try:
                error_df = pd.DataFrame(self.error_records)
                error_df.to_excel(file_path, index=False)
                QMessageBox.information(self, "成功", f"错误记录导出成功！\n共 {len(self.error_records)} 条记录")
            except Exception as e:
                QMessageBox.critical(self, "错误", f"导出失败：{str(e)}")
